"""
online_b2b.services.flipkart_fsn
================================

Fill a Flipkart line's **blank EAN from its FSN**, and say so out loud.

Why this exists: Flipkart's PO export sometimes ships a line with the EAN cell
EMPTY while the FSN is present. The frozen dump parser treats a row whose EAN
isn't 8–14 digits as a footer row and ``continue``s past it, so such a line was
**silently dropped** — on 07-09-2026 that quietly lost 504 units of one lip
liner across four POs, with nothing on screen to show for it.

The engine is frozen, so the fix runs BEFORE it: we rewrite a temp copy of the
workbook with the EAN filled in from the FSN, exactly as if Flipkart had sent it
complete. Same pattern as the MT ``_normalize_*_excel`` pre-passes.

The FSN→EAN map is DATA, not code — ``channel_sku_map`` rows with
``channel='Flipkart'`` (the table Swiggy already uses for SkuCode→EAN). A new
FSN is one row, no deploy.

Nothing is ever guessed: an FSN we don't know is reported by name, quantity and
description so the operator can map it, rather than vanishing.
"""
from __future__ import annotations

import logging
import os
import tempfile

log = logging.getLogger(__name__)
CHANNEL = 'Flipkart'


def fsn_map() -> dict:
    """``{FSN: (ean, item_no)}`` for Flipkart, from ``channel_sku_map``."""
    out = {}
    try:
        from .order_db import _conn
        with _conn() as (cur, d):
            cur.execute(
                f"SELECT c.sku_code, COALESCE(NULLIF(m.ean,''), c.ean), c.item_no "
                f"FROM channel_sku_map c "
                f"LEFT JOIN item_master m ON m.item_no = c.item_no "
                f"WHERE c.channel = {d['ph']}", (CHANNEL,))
            for code, ean, item_no in cur.fetchall():
                key = str(code or '').strip().upper()
                if key:
                    out[key] = (str(ean or '').strip(), str(item_no or '').strip())
    except Exception:  # noqa: BLE001 — no DB → empty map, callers still warn
        log.exception('Flipkart FSN map unavailable')
    return out


def fill_blank_eans(paths):
    """Rewrite any Flipkart workbook that has blank-EAN rows. -> (paths, notes).

    Returns the paths to hand the engine (a temp copy where something was
    filled, the original otherwise) and human-readable notes for the review page.
    """
    notes: list[str] = []
    try:
        import openpyxl
        import pandas as pd
        from online_po_processor.engine.flipkart_dump_parser import (
            _build_col_map, _header_top_row, _norm)
    except Exception as e:  # noqa: BLE001 — never break the upload
        log.warning('flipkart_fsn unavailable: %s', e)
        return list(paths), notes

    mapping = fsn_map()
    out_paths, filled_total, unknown_total = [], 0, 0

    for path in paths:
        if not str(path).lower().endswith(('.xlsx', '.xlsm')):
            out_paths.append(path)
            continue
        try:
            raw = pd.read_excel(path, header=None, dtype=str).fillna('')
            top = _header_top_row(raw)
            cmap = _build_col_map(raw, top) if top is not None else {}
        except Exception:  # noqa: BLE001 — not a dump we understand; leave it
            out_paths.append(path)
            continue
        if top is None or 'ean' not in cmap or 'fsn' not in cmap:
            out_paths.append(path)
            continue

        i_ean, i_fsn = cmap['ean'], cmap['fsn']
        i_qty, i_desc = cmap.get('qty'), cmap.get('description')
        edits = []                                  # (row_index, ean)
        for i in range(top + 2, len(raw)):
            cells = raw.iloc[i].tolist()
            c0 = _norm(cells[0]) if cells else ''
            if c0.startswith('totalquantity') or c0.startswith('total='):
                break
            def cell(ci):
                return str(cells[ci]).strip() if (ci is not None and ci < len(cells)) else ''
            if cell(i_ean):                          # EAN present — nothing to do
                continue
            fsn = cell(i_fsn)
            if not fsn:                              # blank row / spacer
                continue
            # The footer ('Total Quantity= 1392') sits in the FSN column on some
            # exports, so the parser's column-0 break never fires for it. Skip it
            # here too — reporting a footer as an "unmapped FSN" is noise, and
            # noisy warnings are how real ones get ignored.
            if _norm(fsn).startswith(('total', 'grandtotal')):
                continue
            qty, desc = cell(i_qty), cell(i_desc)[:48]
            hit = mapping.get(fsn.upper())
            if hit and hit[0]:
                edits.append((i, hit[0]))
                notes.append(
                    f"Flipkart {os.path.basename(str(path))}: FSN {fsn} had a BLANK "
                    f"EAN — filled {hit[0]}"
                    + (f" (item {hit[1]})" if hit[1] else '')
                    + f" from the FSN map. Qty {qty or '?'} · {desc}")
            else:
                unknown_total += 1
                notes.append(
                    f"Flipkart {os.path.basename(str(path))}: FSN {fsn} has a BLANK "
                    f"EAN and is NOT in the FSN map — this line will be DROPPED "
                    f"(qty {qty or '?'} · {desc}). Add it under Item Master → "
                    f"Channel SKU Map (channel 'Flipkart') to include it.")

        if not edits:
            out_paths.append(path)
            continue
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb[wb.sheetnames[0]]
            for row_i, ean in edits:
                ws.cell(row=row_i + 1, column=i_ean + 1, value=ean)   # 1-based
            # Preserve the ORIGINAL filename — Flipkart derives its PO number
            # from the file's basename (purchase_order_<PO>.xlsx). A random
            # mkstemp name (tmpXXXX_fk_fsn.xlsx) made the parser read the PO as
            # 'TMPXXXX_FK_FSN'. Write into a fresh temp DIR under the real name.
            tmpdir = tempfile.mkdtemp(prefix='fk_fsn_')
            tmp = os.path.join(tmpdir, os.path.basename(str(path)))
            wb.save(tmp)
            out_paths.append(tmp)
            filled_total += len(edits)
        except Exception as e:  # noqa: BLE001 — fall back to the original file
            log.warning('could not rewrite %s: %s', path, e)
            notes.append(f"Flipkart {os.path.basename(str(path))}: could not fill "
                         f"the blank EAN(s) ({e}) — line(s) will be dropped.")
            out_paths.append(path)

    if filled_total or unknown_total:
        notes.insert(0, f"Flipkart FSN check: {filled_total} blank EAN(s) filled "
                        f"from the FSN map, {unknown_total} unmapped.")
    return out_paths, notes
